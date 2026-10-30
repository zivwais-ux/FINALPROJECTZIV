"""Break the Crew 1 -> Crew 2 handoff on purpose and show that the flow stops.

Each scenario copies Crew 1's outputs to artifacts/broken/<scenario>/, breaks
the copy, and runs the flow on it. artifacts/latest/ is never touched.

Handoff scenarios run with --from-gate (no LLM, a few seconds each): the data
changes after Crew 1 wrote the contract, as in the Harbor & Vale incident.
raw_unit_change breaks a copy of the raw Excel file and runs the whole flow,
Crew 1 included (needs OPENAI_API_KEY): Crew 1 then writes a contract that
matches the broken data, and only the locked contract catches it.

    uv run python scripts/break_it.py                  # the five handoff scenarios
    uv run python scripts/break_it.py unit_change      # one scenario
    uv run python scripts/break_it.py raw_unit_change  # full run with the LLM

Exits with 0 when the flow caught every scenario it ran, 1 otherwise.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd

from churn_pipeline import main as pipeline
from churn_pipeline.scenarios import SCENARIOS, Scenario

ROOT = Path(__file__).resolve().parents[1]
CREW1 = ROOT / "artifacts" / "latest" / "crew1"
RAW = ROOT / "data" / "raw" / "E Commerce Dataset.xlsx"
BROKEN = ROOT / "artifacts" / "broken"


RAW_SCENARIO = Scenario(
    "CashbackAmount x100 in the raw Excel file; Crew 1 rebuilds the contract on it",
    lambda df: df, "locked_contract", "columns.CashbackAmount.expected_mean")


def run_scenario(name: str, source: Path = CREW1, out_root: Path = BROKEN) -> dict:
    """Break a copy of the handoff (or of the raw file) and run the flow on it."""
    out = out_root / name
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True)
    inputs = {"output_dir": str(out / "crew1"), "crew2_output_dir": str(out / "crew2"),
              "report_path": str(out / "validation_report.md"), "run_kind": "scenario", "scenario": name}

    if name == "raw_unit_change":
        scenario = RAW_SCENARIO
        raw = out / "raw" / RAW.name
        raw.parent.mkdir()
        wb = openpyxl.load_workbook(RAW)
        ws = wb["E Comm"]
        col = [c.value for c in ws[1]].index("CashbackAmount") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col)
            if isinstance(cell.value, (int, float)):
                cell.value = round(cell.value * 100, 2)
        wb.save(raw)
        inputs["raw_data_path"] = str(raw)
    else:
        scenario = SCENARIOS[name]
        shutil.copytree(source, out / "crew1")
        data = out / "crew1" / "clean_data.csv"
        scenario.change(pd.read_csv(data)).to_csv(data, index=False)
        inputs["from_gate"] = True

    pipeline._setup_logging(out / "pipeline.log")
    pipeline.log.info("break_it scenario %s: %s", name, scenario.description)
    result = pipeline.ChurnPipelineFlow().kickoff(inputs=inputs)

    failed = {g["name"]: g["problems"] for g in result["gates"] if not g["passed"]}
    proof = next((p for p in failed.get(scenario.expected_gate, []) if scenario.expected_text in p), "")
    caught = result["status"] != "passed" and bool(proof)
    return {
        "scenario": name,
        "description": scenario.description,
        "status": result["status"],
        "failed_gates": list(failed),
        "caught": caught,
        "crew2_ran": (out / "crew2").exists(),
        "problem": proof,
        "report": str(out / "validation_report.md"),
    }


def write_summary(results: list[dict], path: Path) -> None:
    lines = ["# break_it results", "",
             "| Scenario | Break | Caught | Failed gates | Crew 2 ran | Problem reported |",
             "|---|---|---|---|---|---|"]
    for r in results:
        lines.append(f"| `{r['scenario']}` | {r['description']} | {'yes' if r['caught'] else 'NO'} | "
                     f"{', '.join(r['failed_gates']) or 'none'} | {'yes' if r['crew2_ran'] else 'no'} | "
                     f"{r['problem'].replace('|', '/')} |")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: list[str]) -> int:
    names = argv or list(SCENARIOS)
    unknown = [n for n in names if n not in SCENARIOS and n != "raw_unit_change"]
    if unknown:
        print(f"unknown scenario(s): {unknown}; choose from {list(SCENARIOS) + ['raw_unit_change']}")
        return 2
    results = [run_scenario(n) for n in names]
    summary = BROKEN / ("summary.md" if not argv else f"summary_{'_'.join(names)}.md")
    write_summary(results, summary)
    print(summary.read_text(encoding="utf-8"))
    return 0 if all(r["caught"] and not r["crew2_ran"] for r in results) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
